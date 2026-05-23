#!/usr/bin/env python3
"""Build a cross-tab xlsx of round1 vs overall metrics across benchmarks."""
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[3]
OUTPUT = REPO / "benchmark" / "results" / "summary_table.xlsx"

METRICS = [
    "compile_pass_rate",
    "correct_pass_rate",
    "positive_rate",
    "geometric_mean_speedup",
    "max_speedup",
]
SECTIONS = [("round1", "首次迭代"), ("overall", "最终")]

BENCHMARKS = ["tsvc", "polybenchc", "rajaperf"]
MODELS = [
    "Llama3.1-8B",
    "Llama3.1-8B-FT",
    "Qwen3.5-9B",
    "Qwen3.5-9B-FT",
    "Qwen3.5-27B",
    "Qwen3.5-27B-FT",
    "DeepSeek-V4-Pro",
    "Kimi-K2.5",
    "gpt5.4",
]


def model_dir_candidates(model: str):
    """Map a display name to plausible on-disk directory names."""
    yield model
    yield model.lower()
    yield model.lower().replace(".", "_")
    yield model.replace("-", "_").lower()


def find_results_path(benchmark: str, model: str) -> Path | None:
    base = REPO / "benchmark" / "results" / benchmark
    if not base.exists():
        return None
    for cand in model_dir_candidates(model):
        path = base / cand / "results.jsonl"
        if path.exists():
            return path
    return None


ROWS = [(b, m, find_results_path(b, m)) for b in BENCHMARKS for m in MODELS]


def load_summary(jsonl_path: Path | None) -> dict | None:
    if jsonl_path is None:
        return None
    summary_path = jsonl_path.parent / "summary.json"
    if not summary_path.exists():
        return None
    with summary_path.open(encoding="utf-8") as f:
        return json.load(f)


def fmt(value):
    if value is None:
        return ""
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "summary"

    thin = Side(border_style="thin", color="888888")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    section_fill = PatternFill("solid", fgColor="B8CCE4")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Row 1: top section labels (merged across the 5 metric columns each)
    ws.cell(row=1, column=1, value="benchmark")
    ws.cell(row=1, column=2, value="model")
    for idx, (_, label) in enumerate(SECTIONS):
        first = 3 + idx * len(METRICS)
        last = first + len(METRICS) - 1
        ws.cell(row=1, column=first, value=label)
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)

    # Row 2: metric column headers
    ws.cell(row=2, column=1, value="")
    ws.cell(row=2, column=2, value="")
    for sec_idx, _ in enumerate(SECTIONS):
        for m_idx, metric in enumerate(METRICS):
            ws.cell(row=2, column=3 + sec_idx * len(METRICS) + m_idx, value=metric)

    # Merge benchmark/model labels vertically
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    # Style headers
    for col in range(1, 3 + len(SECTIONS) * len(METRICS)):
        for row in (1, 2):
            cell = ws.cell(row=row, column=col)
            cell.alignment = center
            cell.font = bold
            cell.border = border
            cell.fill = section_fill if row == 1 else header_fill

    # Data rows
    for r_idx, (bench, model, path) in enumerate(ROWS, start=3):
        ws.cell(row=r_idx, column=1, value=bench).alignment = center
        ws.cell(row=r_idx, column=2, value=model).alignment = center
        summary = load_summary(path)
        for sec_idx, (sec_key, _) in enumerate(SECTIONS):
            block = (summary or {}).get(sec_key, {}) if summary else {}
            for m_idx, metric in enumerate(METRICS):
                col = 3 + sec_idx * len(METRICS) + m_idx
                cell = ws.cell(row=r_idx, column=col, value=fmt(block.get(metric)))
                cell.number_format = "0.0000"
                cell.alignment = center
                cell.border = border
        for col in range(1, 3 + len(SECTIONS) * len(METRICS)):
            ws.cell(row=r_idx, column=col).border = border

    # Column widths
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 11
    for col in range(3, 3 + len(SECTIONS) * len(METRICS)):
        ws.column_dimensions[get_column_letter(col)].width = 16

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 38

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
